"""
Reward Script: Create a bar chart from pivot table data on a new chart sheet
Task ID: calc_gg5_041
Domain: libreoffice_calc
Scoring:
  Component 1: HoursChart sheet exists (0.20 pts)
  Component 2: Bar/column chart on HoursChart sheet (0.25 pts)
  Component 3: Chart data sourced from Summary sheet pivot table (0.20 pts)
  Component 4: Data labels enabled showing values (0.20 pts)
  Component 5: Chart has title and axis labels (0.15 pts)
"""

import os

WORKDIR = '/home/user'
TASK_ID = 'calc_gg5_041'


def persist_app_state(domain):
    """Save any unsaved LibreOffice changes before verification."""
    import time
    os.environ["DISPLAY"] = ":0"
    if domain in {"libreoffice_calc", "libreoffice_writer", "libreoffice_impress"}:
        try:
            import pyautogui
            pyautogui.hotkey("ctrl", "s")
            time.sleep(0.8)
            print(f"PERSIST: ctrl+s sent for {domain}")
        except Exception as e:
            print(f"PERSIST_WARN: save hook failed: {e}")


def verify_task(file_path):
    """
    Verify task completion with progressive scoring.
    Returns: float between 0.0 and 1.0
    """
    import openpyxl

    total_score = 0.0

    try:
        wb = openpyxl.load_workbook(file_path)
    except Exception as e:
        print(f"CRITICAL: Cannot load file {file_path}: {e}")
        print("REWARD: 0.0")
        return 0.0

    sheet_names = wb.sheetnames
    print(f"INFO: Sheet names: {sheet_names}")

    # Component 1: HoursChart sheet exists (0.20 points)
    # This sheet does NOT exist in initial_env, only in golden_env
    try:
        if 'HoursChart' in sheet_names:
            print(f"PASS: Component 1 — 'HoursChart' sheet exists (0.20 pts)")
            total_score += 0.20
        else:
            print(f"FAIL: Component 1 — 'HoursChart' sheet not found in {sheet_names}")
    except Exception as e:
        print(f"ERROR: Component 1 — {e}")

    # Component 2: Bar/column chart exists on HoursChart sheet (0.25 points)
    # No charts exist in initial_env at all
    chart = None
    try:
        if 'HoursChart' in sheet_names:
            cs = wb['HoursChart']
            charts = cs._charts if hasattr(cs, '_charts') else []
            if len(charts) >= 1:
                chart = charts[0]
                chart_type = chart.type
                # Accept both 'col' (vertical bar) and 'bar' (horizontal bar) as bar charts
                if chart_type in ('col', 'bar'):
                    print(f"PASS: Component 2 — Bar/column chart found on HoursChart (type={chart_type}) (0.25 pts)")
                    total_score += 0.25
                else:
                    print(f"FAIL: Component 2 — Chart exists but type is '{chart_type}', expected 'col' or 'bar'")
            else:
                print(f"FAIL: Component 2 — No charts found on HoursChart sheet")
        else:
            print(f"FAIL: Component 2 — HoursChart sheet does not exist")
    except Exception as e:
        print(f"ERROR: Component 2 — {e}")

    # Component 3: Chart data references Summary sheet employee names (col A) and total hours (col F) (0.20 points)
    # Initial_env has no chart, so this naturally fails on initial
    try:
        if chart is not None and len(chart.series) >= 1:
            series = chart.series[0]
            val_ref = None
            cat_ref = None

            # Extract value reference
            if hasattr(series, 'val') and series.val and series.val.numRef:
                val_ref = series.val.numRef.f
            # Extract category reference
            if hasattr(series, 'cat') and series.cat:
                if series.cat.strRef:
                    cat_ref = series.cat.strRef.f
                elif series.cat.numRef:
                    cat_ref = series.cat.numRef.f

            print(f"INFO: val_ref={val_ref}, cat_ref={cat_ref}")

            # Check that data comes from Summary sheet
            refs_summary = False
            if val_ref and 'Summary' in val_ref:
                refs_summary = True
            if cat_ref and 'Summary' in cat_ref:
                refs_summary = True

            # Check that value data references column F (Total Hours)
            val_col_f = False
            if val_ref and '$F$' in val_ref.upper():
                val_col_f = True

            if refs_summary and val_col_f:
                print(f"PASS: Component 3 — Chart data sourced from Summary sheet column F (0.20 pts)")
                total_score += 0.20
            elif refs_summary:
                print(f"PARTIAL: Component 3 — References Summary but not specifically column F (val_ref={val_ref})")
                total_score += 0.10
            else:
                print(f"FAIL: Component 3 — Chart data does not reference Summary sheet (val={val_ref}, cat={cat_ref})")
        else:
            print(f"FAIL: Component 3 — No chart or no series to check")
    except Exception as e:
        print(f"ERROR: Component 3 — {e}")

    # Component 4: Data labels enabled showing values (0.20 points)
    # Initial_env has no chart, so this naturally fails on initial
    try:
        if chart is not None:
            data_labels_show_val = False

            # Check chart-level data labels
            if hasattr(chart, 'dataLabels') and chart.dataLabels:
                if chart.dataLabels.showVal:
                    data_labels_show_val = True

            # Also check series-level data labels
            if not data_labels_show_val:
                for s in chart.series:
                    if hasattr(s, 'dLbls') and s.dLbls and s.dLbls.showVal:
                        data_labels_show_val = True
                        break

            if data_labels_show_val:
                print(f"PASS: Component 4 — Data labels with showVal=True enabled (0.20 pts)")
                total_score += 0.20
            else:
                print(f"FAIL: Component 4 — Data labels not showing values")
        else:
            print(f"FAIL: Component 4 — No chart to check data labels")
    except Exception as e:
        print(f"ERROR: Component 4 — {e}")

    # Component 5: Chart has a meaningful title and axis labels (0.15 points)
    # Initial_env has no chart, so this naturally fails on initial
    try:
        if chart is not None:
            sub_score = 0.0

            # Check chart title exists
            has_title = False
            if chart.title:
                # Try to extract title text
                try:
                    if chart.title.tx and chart.title.tx.rich:
                        for p in chart.title.tx.rich.paragraphs:
                            for r in p.r:
                                if r.t and len(r.t.strip()) > 0:
                                    has_title = True
                                    print(f"INFO: Chart title text: '{r.t}'")
                except Exception:
                    pass
                if not has_title:
                    # Title object exists even without rich text
                    has_title = True

            if has_title:
                sub_score += 0.05
                print(f"PASS: Component 5a — Chart has a title")
            else:
                print(f"FAIL: Component 5a — Chart has no title")

            # Check y-axis title
            has_y_title = False
            if chart.y_axis and chart.y_axis.title:
                has_y_title = True
                sub_score += 0.05
                print(f"PASS: Component 5b — Y-axis has a title")
            else:
                print(f"FAIL: Component 5b — Y-axis has no title")

            # Check x-axis title
            has_x_title = False
            if chart.x_axis and chart.x_axis.title:
                has_x_title = True
                sub_score += 0.05
                print(f"PASS: Component 5c — X-axis has a title")
            else:
                print(f"FAIL: Component 5c — X-axis has no title")

            if sub_score > 0:
                print(f"PASS: Component 5 — Chart titles/labels ({sub_score:.2f} pts)")
                total_score += sub_score
            else:
                print(f"FAIL: Component 5 — No chart title or axis labels found")
        else:
            print(f"FAIL: Component 5 — No chart to check titles")
    except Exception as e:
        print(f"ERROR: Component 5 — {e}")

    final_score = round(min(total_score, 1.0), 2)
    print(f"\nScore: {final_score}/1.0")
    print(f"REWARD: {final_score}")
    return final_score


# Entry point
persist_app_state("libreoffice_calc")

file_path = f'{WORKDIR}/{TASK_ID}.xlsx'
if not os.path.exists(file_path):
    print(f"File not found: {file_path}")
    print("REWARD: 0.0")
else:
    verify_task(file_path)
