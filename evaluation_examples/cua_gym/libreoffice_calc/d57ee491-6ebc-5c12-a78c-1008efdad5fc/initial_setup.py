"""
Initial Setup: Multi-app task — Program synthesis paper open in PDF viewer,
               blank Writer document open, Chrome available.
Task ID: osworld_multi_apps_paper_scholar_browse_015
Domain: multi_apps (libreoffice_writer + chrome + pdf)
"""

import os
import shlex
import subprocess
import time

WORKDIR = '/home/user'
TASK_ID = 'osworld_multi_apps_paper_scholar_browse_015'
PDF_OUTPUT = f'{WORKDIR}/{TASK_ID}.pdf'
DOCX_OUTPUT = f'{WORKDIR}/{TASK_ID}.docx'


def launch_gui(command: str, delay_sec: float = 1.0):
    """Launch GUI app on VM display without blocking script exit."""
    env = os.environ.copy()
    env["DISPLAY"] = ":0"
    subprocess.Popen(
        shlex.split(command),
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        env=env,
    )
    time.sleep(delay_sec)


def create_pdf():
    """Create a realistic program synthesis paper PDF with Armando Solar-Lezama as first author."""
    try:
        from fpdf import FPDF
    except ImportError:
        subprocess.run(['pip3', 'install', 'fpdf2'], check=True)
        from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    # Title
    pdf.set_font('Helvetica', 'B', 16)
    pdf.set_xy(20, 20)
    pdf.multi_cell(170, 8, 'Sketching Program Structures with Angelic Nondeterminism', align='C')

    # Authors
    pdf.set_font('Helvetica', '', 12)
    pdf.set_xy(20, 42)
    pdf.multi_cell(170, 6,
        'Armando Solar-Lezama, Rodric Rabbah, Rastislav Bodik, Kemal Ebcioglu',
        align='C')

    # Affiliations
    pdf.set_font('Helvetica', 'I', 10)
    pdf.set_xy(20, 56)
    pdf.multi_cell(170, 5,
        'Massachusetts Institute of Technology, Cambridge, MA 02139\n'
        'IBM T.J. Watson Research Center, Yorktown Heights, NY 10598\n'
        'University of California, Berkeley, CA 94720',
        align='C')

    # Separator line simulation
    pdf.set_xy(20, 76)
    pdf.set_font('Helvetica', '', 10)

    # Abstract
    pdf.set_font('Helvetica', 'B', 11)
    pdf.set_xy(20, 80)
    pdf.cell(0, 6, 'Abstract')

    pdf.set_font('Helvetica', '', 10)
    pdf.set_xy(20, 88)
    pdf.multi_cell(170, 5,
        'We present a new approach to program synthesis that combines sketching with '
        'angelic nondeterminism. Our technique allows programmers to write partial '
        'programs called sketches, which are then completed automatically by a synthesis '
        'engine. The synthesis engine uses a combination of counterexample-guided inductive '
        'synthesis (CEGIS) and SAT-based verification to find complete implementations '
        'that satisfy the specification. We demonstrate the effectiveness of our approach '
        'on a range of benchmarks including data structure manipulations, bit-manipulation '
        'routines, and stencil computations. Our results show significant speedups compared '
        'to previous synthesis approaches while maintaining correctness guarantees.')

    # 1. Introduction
    pdf.set_font('Helvetica', 'B', 11)
    pdf.set_xy(20, 148)
    pdf.cell(0, 6, '1. Introduction')

    pdf.set_font('Helvetica', '', 10)
    pdf.set_xy(20, 156)
    pdf.multi_cell(170, 5,
        'Program synthesis is the task of automatically constructing a program that '
        'satisfies a given specification. While early work in program synthesis focused '
        'on complete specifications given as formal logic, recent approaches have explored '
        'more practical specification mechanisms including input-output examples, reference '
        'implementations, and partial programs known as sketches [Solar-Lezama et al. 2006].\n\n'
        'The sketch-based synthesis approach, pioneered by Solar-Lezama, allows a programmer '
        'to provide the high-level structure of a program while leaving low-level details '
        'to be filled in by the synthesizer. This dramatically reduces the search space '
        'compared to synthesis from scratch, making the approach practical for real programs.\n\n'
        'In this paper, we extend the sketch-based approach with angelic nondeterminism, '
        'a technique that allows the synthesis engine to explore multiple possible completions '
        'simultaneously. This enables more efficient search and better scalability to larger '
        'synthesis problems.')

    # 2. Background
    pdf.add_page()
    pdf.set_font('Helvetica', 'B', 11)
    pdf.set_xy(20, 20)
    pdf.cell(0, 6, '2. Background and Related Work')

    pdf.set_font('Helvetica', '', 10)
    pdf.set_xy(20, 28)
    pdf.multi_cell(170, 5,
        'Program synthesis has a rich history dating back to the 1970s, but has seen '
        'a renaissance in recent years due to advances in constraint solvers and '
        'verification techniques. Key milestones include:\n\n'
        '- Inductive Logic Programming (ILP): Learning programs from examples using '
        'first-order logic representations.\n\n'
        '- Counterexample-Guided Inductive Synthesis (CEGIS): An iterative approach '
        'that alternates between finding candidate programs and verifying their correctness.\n\n'
        '- Sketch-based synthesis: Allowing programmers to provide structural hints '
        'to guide the synthesis process.\n\n'
        '- Neural program synthesis: Using deep learning to generate programs from '
        'natural language descriptions or input-output examples.\n\n'
        'Our work builds on CEGIS and sketch-based synthesis, extending both with '
        'angelic nondeterminism to improve scalability and performance.')

    # 3. Approach
    pdf.set_font('Helvetica', 'B', 11)
    pdf.set_xy(20, 118)
    pdf.cell(0, 6, '3. The Angelic Synthesis Approach')

    pdf.set_font('Helvetica', '', 10)
    pdf.set_xy(20, 126)
    pdf.multi_cell(170, 5,
        'Our approach extends sketch-based synthesis with angelic nondeterminism. '
        'The key insight is that instead of searching for a single completion of the '
        'sketch, we allow the synthesis engine to maintain multiple possible completions '
        'simultaneously, pruning them based on counterexamples.\n\n'
        'Formally, given a sketch S and a specification phi, our goal is to find '
        'a complete program P such that P satisfies phi for all inputs. The angelic '
        'synthesis algorithm proceeds as follows:\n\n'
        '1. Initialize the set of candidate completions C = {all possible completions of S}\n'
        '2. While C is not empty:\n'
        '   a. Select a candidate P from C\n'
        '   b. Verify P against phi\n'
        '   c. If P satisfies phi, return P\n'
        '   d. Otherwise, add the counterexample to the constraint set and prune C\n'
        '3. Return UNSAT if no valid completion exists')

    # References
    pdf.add_page()
    pdf.set_font('Helvetica', 'B', 11)
    pdf.set_xy(20, 20)
    pdf.cell(0, 6, 'References')

    pdf.set_font('Helvetica', '', 9)
    pdf.set_xy(20, 28)
    pdf.multi_cell(170, 5,
        '[1] Solar-Lezama, A., Tancau, L., Bodik, R., Seshia, S., and Saraswat, V. '
        'Combinatorial sketching for finite programs. In ASPLOS 2006.\n\n'
        '[2] Solar-Lezama, A. Program synthesis by sketching. PhD thesis, UC Berkeley, 2008.\n\n'
        '[3] Gulwani, S. Automating string processing in spreadsheets using input-output '
        'examples. In POPL 2011.\n\n'
        '[4] Jha, S., Gulwani, S., Seshia, S. A., and Tiwari, A. Oracle-guided '
        'component-based program synthesis. In ICSE 2010.\n\n'
        '[5] Alur, R., Bodik, R., Juniwal, G., Martin, M. M. K., Raghothaman, M., '
        'Seshia, S. A., Singh, R., Solar-Lezama, A., Torlak, E., and Udupa, A. '
        'Syntax-guided synthesis. In FMCAD 2013.')

    pdf.output(PDF_OUTPUT)
    print(f'PDF created: {PDF_OUTPUT}')


def create_blank_docx():
    """Create a blank LibreOffice Writer document."""
    try:
        from docx import Document
    except ImportError:
        subprocess.run(['pip3', 'install', 'python-docx'], check=True)
        from docx import Document

    doc = Document()
    # Remove default empty paragraph content but keep document structure
    # The document should be blank (empty) — no pre-written content
    # Clear all paragraphs
    for para in doc.paragraphs:
        para.clear()

    doc.save(DOCX_OUTPUT)
    print(f'Blank Writer document created: {DOCX_OUTPUT}')


def main():
    create_pdf()
    create_blank_docx()

    # Kill any existing Chrome instances before launching
    subprocess.run(['pkill', '-f', 'chrome'], capture_output=True)
    time.sleep(1)

    # Launch the PDF in Evince (default PDF viewer on Ubuntu)
    launch_gui(f'evince "{PDF_OUTPUT}"', delay_sec=2.0)

    # Launch LibreOffice Writer with the blank document
    launch_gui(f'libreoffice --writer "{DOCX_OUTPUT}"', delay_sec=2.0)

    # Launch Chrome (blank / new tab — agent will navigate to Scholar)
    launch_gui('google-chrome --new-window', delay_sec=2.0)

    print('GUI_READY: launched evince (PDF), LibreOffice Writer (blank doc), Chrome')


main()
