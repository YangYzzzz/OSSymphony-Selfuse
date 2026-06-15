"""
Reward Script: Create a Gantt-style stacked bar chart for project timeline
Task ID: calc_gcp_047
Domain: libreoffice_calc
Scoring:
  Component 1 (0.15): Chart exists on ProjectPlan sheet
  Component 2 (0.20): Chart is horizontal stacked bar (type=bar, grouping=stacked)
  Component 3 (0.15): Chart has exactly 2 series (StartDay invisible + Duration colored)
  Component 4 (0.25): First series (StartDay offset) is invisible (noFill)
  Component 5 (0.10): Second series (Duration) has visible fill (colored)
  Component 6 (0.15): Chart has meaningful title referencing Gantt/Timeline/Project
"""

import os
import openpyxl

WORKDIR = '/home/user'
TASK_ID = 'calc_gcp_047'


def persist_app_state(domain: str):
    """Save any unsaved GUI edits before verification."""
    import time
    os.environ["DISPLAY"] = ":0"
    if domain in {"libreoffice_calc", "libreoffice_writer", "libreoffice_impress"}:
        try:
            import pyautogui
            pyautogui.hotkey("ctrl", "s")
            time.sleep(0.8)
            print(f"PERSIST: ctrl+s sent for {domain}")
        except Exception as e:
            print(f"PERSIST_WARN: save hook failed: {e}")


def verify_task(file_path):
    """
    Verify Gantt-style stacked bar chart creation with progressive scoring.
    Returns: float between 0.0 and 1.0
    """
    total_score = 0.0

    # Load workbook
    try:
        wb = openpyxl.load_workbook(file_path)
    except Exception as e:
        print(f"CRITICAL: Cannot load file {file_path}: {e}")
        print("REWARD: 0.0")
        return 0.0

    # Check that ProjectPlan sheet exists (precondition gate)
    if 'ProjectPlan' not in wb.sheetnames:
        print("CRITICAL: 'ProjectPlan' sheet not found")
        print("REWARD: 0.0")
        return 0.0

    ws = wb['ProjectPlan']

    # Component 1: Chart exists on ProjectPlan sheet (0.15 points)
    # Initial has 0 charts, golden has 1+
    try:
        charts = ws._charts
        num_charts = len(charts)
        if num_charts >= 1:
            print(f"PASS: Component 1 -- Chart exists on ProjectPlan ({num_charts} chart(s)) (0.15 pts)")
            total_score += 0.15
        else:
            print(f"FAIL: Component 1 -- No charts found on ProjectPlan sheet")
    except Exception as e:
        print(f"ERROR: Component 1 -- {e}")

    # If no charts, remaining checks are moot
    if total_score < 0.1:
        final_score = min(total_score, 1.0)
        print(f"\nScore: {total_score}/1.0")
        print(f"REWARD: {final_score}")
        return final_score

    chart = charts[0]

    # Component 2: Chart is horizontal stacked bar (0.20 points)
    # type="bar" (horizontal bars) and grouping="stacked"
    try:
        is_bar_chart = type(chart).__name__ == 'BarChart'
        chart_type = getattr(chart, 'type', None)
        chart_grouping = getattr(chart, 'grouping', None)

        type_ok = is_bar_chart and chart_type == 'bar'
        grouping_ok = chart_grouping == 'stacked'

        if type_ok and grouping_ok:
            print(f"PASS: Component 2 -- Horizontal stacked bar chart (type={chart_type}, grouping={chart_grouping}) (0.20 pts)")
            total_score += 0.20
        else:
            # Partial: if it's a bar chart but wrong subtype or grouping
            if is_bar_chart and grouping_ok:
                # stacked but col instead of bar -- still partially correct
                print(f"PARTIAL: Component 2 -- Stacked but type={chart_type} (expected 'bar') (0.10 pts)")
                total_score += 0.10
            elif is_bar_chart and type_ok:
                print(f"PARTIAL: Component 2 -- Horizontal bar but grouping={chart_grouping} (expected 'stacked') (0.10 pts)")
                total_score += 0.10
            else:
                print(f"FAIL: Component 2 -- Expected horizontal stacked bar, found {type(chart).__name__} type={chart_type} grouping={chart_grouping}")
    except Exception as e:
        print(f"ERROR: Component 2 -- {e}")

    # Component 3: Chart has exactly 2 series (0.15 points)
    # Series 0 = StartDay (invisible offset), Series 1 = Duration (colored)
    try:
        num_series = len(chart.series)
        if num_series == 2:
            print(f"PASS: Component 3 -- Chart has 2 series (0.15 pts)")
            total_score += 0.15
        elif num_series >= 2:
            # More than 2 series but at least has the right structure
            print(f"PARTIAL: Component 3 -- Chart has {num_series} series (expected 2) (0.07 pts)")
            total_score += 0.07
        else:
            print(f"FAIL: Component 3 -- Chart has {num_series} series (expected 2)")
    except Exception as e:
        print(f"ERROR: Component 3 -- {e}")

    # Component 4: First series (StartDay offset) is invisible (0.25 points)
    # The first series should have noFill (both fill and line) to create
    # the invisible offset that makes the Gantt appearance
    try:
        if len(chart.series) >= 1:
            s0 = chart.series[0]
            gp = s0.graphicalProperties

            fill_invisible = (gp is not None and bool(gp.noFill))
            line_invisible = (gp is not None and hasattr(gp, 'line')
                              and gp.line is not None and bool(gp.line.noFill))

            if fill_invisible and line_invisible:
                print(f"PASS: Component 4 -- First series is fully invisible (noFill + line noFill) (0.25 pts)")
                total_score += 0.25
            elif fill_invisible:
                print(f"PARTIAL: Component 4 -- First series fill is invisible but line may be visible (0.15 pts)")
                total_score += 0.15
            else:
                print(f"FAIL: Component 4 -- First series is NOT invisible (noFill={getattr(gp, 'noFill', None)})")
        else:
            print(f"FAIL: Component 4 -- No series to check")
    except Exception as e:
        print(f"ERROR: Component 4 -- {e}")

    # Component 5: Second series (Duration) has visible colored fill (0.10 points)
    try:
        if len(chart.series) >= 2:
            s1 = chart.series[1]
            gp1 = s1.graphicalProperties

            has_color = False
            if gp1 is not None:
                # Check solidFill exists (explicit color) or default (not noFill)
                if gp1.solidFill is not None or not gp1.noFill:
                    has_color = not bool(gp1.noFill) or gp1.solidFill is not None

            if has_color:
                print(f"PASS: Component 5 -- Second series has visible fill (0.10 pts)")
                total_score += 0.10
            else:
                print(f"FAIL: Component 5 -- Second series has no visible fill")
        else:
            print(f"FAIL: Component 5 -- Less than 2 series")
    except Exception as e:
        print(f"ERROR: Component 5 -- {e}")

    # Component 6: Chart has meaningful title (0.15 points)
    # Should contain Gantt, Timeline, or Project reference
    try:
        title_text = None
        if chart.title is not None:
            # Extract title text
            tx = chart.title.tx
            if tx is not None and tx.rich is not None:
                for para in tx.rich.paragraphs:
                    for run in para.r:
                        if run.t:
                            title_text = (title_text or '') + run.t

        if title_text:
            title_lower = title_text.lower()
            if any(kw in title_lower for kw in ['gantt', 'timeline', 'project']):
                print(f"PASS: Component 6 -- Chart title '{title_text}' references Gantt/Timeline/Project (0.15 pts)")
                total_score += 0.15
            else:
                # Has a title but not a meaningful one for a Gantt chart
                print(f"PARTIAL: Component 6 -- Chart has title '{title_text}' but does not reference Gantt/Timeline/Project (0.07 pts)")
                total_score += 0.07
        else:
            print(f"FAIL: Component 6 -- Chart has no title")
    except Exception as e:
        print(f"ERROR: Component 6 -- {e}")

    final_score = round(min(total_score, 1.0), 2)
    print(f"\nScore: {total_score}/1.0")
    print(f"REWARD: {final_score}")
    return final_score


# Entry point
persist_app_state("libreoffice_calc")

file_path = f'{WORKDIR}/{TASK_ID}.xlsx'
if not os.path.exists(file_path):
    print(f"File not found: {file_path}")
    print("REWARD: 0.0")
else:
    verify_task(file_path)
