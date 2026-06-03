"""
Reward Script: Mail merge fields in release announcement template
Task ID: writer_tech_061
Domain: libreoffice_writer
Scoring:
  Component 1: CustomerName MERGEFIELD exists (0.25 pts)
  Component 2: ProductVersion MERGEFIELD exists (0.25 pts)
  Component 3: ReleaseDate MERGEFIELD exists (0.25 pts)
  Component 4: All placeholder text removed and replaced (0.25 pts)
"""

import os

WORKDIR = '/home/user'
TASK_ID = 'writer_tech_061'


def persist_app_state(domain):
    """Save any unsaved changes in LibreOffice Writer."""
    import time
    os.environ["DISPLAY"] = ":0"
    if domain in {"libreoffice_calc", "libreoffice_writer", "libreoffice_impress"}:
        try:
            import pyautogui
            pyautogui.hotkey("ctrl", "s")
            time.sleep(1.0)
            print("PERSIST: ctrl+s sent for libreoffice_writer")
        except Exception as e:
            print(f"PERSIST_WARN: save hook failed: {e}")


def find_merge_fields(doc):
    """
    Parse the document XML to find all MERGEFIELD instrText entries.
    Returns a dict mapping field name -> count of occurrences.
    """
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    body = doc.element.body
    instr_texts = body.findall('.//w:instrText', ns)

    fields = {}
    for it in instr_texts:
        text = (it.text or '').strip()
        # Expected format: "MERGEFIELD FieldName" possibly with switches
        if text.upper().startswith('MERGEFIELD'):
            parts = text.split()
            if len(parts) >= 2:
                field_name = parts[1]
                fields[field_name] = fields.get(field_name, 0) + 1
    return fields


def verify_task(file_path):
    """
    Verify mail merge field setup with progressive scoring.
    Returns: float between 0.0 and 1.0
    """
    from docx import Document

    total_score = 0.0

    try:
        doc = Document(file_path)
    except Exception as e:
        print(f"CRITICAL: Cannot load file {file_path}: {e}")
        print("REWARD: 0.0")
        return 0.0

    # Extract merge fields from document XML
    merge_fields = find_merge_fields(doc)
    print(f"INFO: Found merge fields: {merge_fields}")

    # Collect all paragraph text for placeholder checks
    all_text = ' '.join(para.text for para in doc.paragraphs)
    print(f"INFO: Document text length: {len(all_text)} chars")

    # Component 1: CustomerName MERGEFIELD exists (0.25 points)
    try:
        if 'CustomerName' in merge_fields:
            print(f"PASS: Component 1 - CustomerName MERGEFIELD found (count={merge_fields['CustomerName']}) (0.25 pts)")
            total_score += 0.25
        else:
            print(f"FAIL: Component 1 - CustomerName MERGEFIELD not found. Fields present: {list(merge_fields.keys())}")
    except Exception as e:
        print(f"ERROR: Component 1 - {e}")

    # Component 2: ProductVersion MERGEFIELD exists (0.25 points)
    try:
        if 'ProductVersion' in merge_fields:
            print(f"PASS: Component 2 - ProductVersion MERGEFIELD found (count={merge_fields['ProductVersion']}) (0.25 pts)")
            total_score += 0.25
        else:
            print(f"FAIL: Component 2 - ProductVersion MERGEFIELD not found. Fields present: {list(merge_fields.keys())}")
    except Exception as e:
        print(f"ERROR: Component 2 - {e}")

    # Component 3: ReleaseDate MERGEFIELD exists (0.25 points)
    try:
        if 'ReleaseDate' in merge_fields:
            print(f"PASS: Component 3 - ReleaseDate MERGEFIELD found (count={merge_fields['ReleaseDate']}) (0.25 pts)")
            total_score += 0.25
        else:
            print(f"FAIL: Component 3 - ReleaseDate MERGEFIELD not found. Fields present: {list(merge_fields.keys())}")
    except Exception as e:
        print(f"ERROR: Component 3 - {e}")

    # Component 4: All original placeholder text removed (0.25 points)
    # The initial file has [Customer Name], [Version], [Release Date] as plain text.
    # After mail merge field insertion, these should be replaced.
    try:
        placeholders_remaining = []
        if '[Customer Name]' in all_text:
            placeholders_remaining.append('[Customer Name]')
        if '[Release Date]' in all_text:
            placeholders_remaining.append('[Release Date]')

        # For [Version], need careful check - it should not appear as a standalone placeholder
        # but "version" as a regular word is fine. Check for exact "[Version]" bracket pattern.
        if '[Version]' in all_text:
            placeholders_remaining.append('[Version]')

        if len(placeholders_remaining) == 0:
            print(f"PASS: Component 4 - All placeholder text removed (0.25 pts)")
            total_score += 0.25
        else:
            print(f"FAIL: Component 4 - Placeholder text still present: {placeholders_remaining}")
    except Exception as e:
        print(f"ERROR: Component 4 - {e}")

    final_score = min(total_score, 1.0)
    print(f"\nScore: {total_score}/1.0")
    print(f"REWARD: {final_score}")
    return final_score


# Entry point
persist_app_state("libreoffice_writer")

file_path = f'{WORKDIR}/{TASK_ID}.docx'
if not os.path.exists(file_path):
    print(f"File not found: {file_path}")
    print("REWARD: 0.0")
else:
    verify_task(file_path)
