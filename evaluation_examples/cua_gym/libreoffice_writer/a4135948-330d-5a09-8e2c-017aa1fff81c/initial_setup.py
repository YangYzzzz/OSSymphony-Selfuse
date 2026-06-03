"""
Initial Setup: Table with 0.5pt solid black borders on all cells
Task ID: writer_tbl_029
Domain: libreoffice_writer
"""

import os
import shlex
import subprocess
import time
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

WORKDIR = '/home/user/Desktop'  # VM path — file lives on Desktop
TASK_ID = 'featured_table'
OUTPUT = f'{WORKDIR}/{TASK_ID}.docx'


def launch_gui(command: str, delay_sec: float = 1.0):
    """Launch GUI app on VM display without blocking script exit."""
    env = os.environ.copy()
    env["DISPLAY"] = ":0"
    subprocess.Popen(
        shlex.split(command),
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        env=env,
    )
    time.sleep(delay_sec)


def make_border_element(tag, val, sz, color, space="0"):
    """Create a border XML element with w:-prefixed attributes."""
    el = OxmlElement(tag)
    el.set(qn('w:val'), val)
    el.set(qn('w:sz'), sz)
    el.set(qn('w:color'), color)
    el.set(qn('w:space'), space)
    return el


def set_table_borders(table, outer_val, outer_sz, outer_color,
                      inner_val="single", inner_sz="4", inner_color="000000"):
    """Set table-level borders via tblPr/tblBorders."""
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)

    # Remove existing tblBorders
    existing = tblPr.find(qn('w:tblBorders'))
    if existing is not None:
        tblPr.remove(existing)

    tblBorders = OxmlElement('w:tblBorders')
    tblBorders.append(make_border_element('w:top', outer_val, outer_sz, outer_color))
    tblBorders.append(make_border_element('w:left', outer_val, outer_sz, outer_color))
    tblBorders.append(make_border_element('w:bottom', outer_val, outer_sz, outer_color))
    tblBorders.append(make_border_element('w:right', outer_val, outer_sz, outer_color))
    tblBorders.append(make_border_element('w:insideH', inner_val, inner_sz, inner_color))
    tblBorders.append(make_border_element('w:insideV', inner_val, inner_sz, inner_color))
    tblPr.append(tblBorders)


def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    """
    Set borders on a table cell using OOXML.
    Each param is a tuple: (val, sz, color) e.g. ("single", "4", "000000")
    """
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()

    existing = tcPr.find(qn('w:tcBorders'))
    if existing is not None:
        tcPr.remove(existing)

    tcBorders = OxmlElement('w:tcBorders')
    for tag, params in [('w:top', top), ('w:left', left), ('w:bottom', bottom), ('w:right', right)]:
        if params is not None:
            val, sz, color = params
            tcBorders.append(make_border_element(tag, val, sz, color))

    tcPr.append(tcBorders)


def create_initial():
    # Ensure Desktop dir exists
    os.makedirs(WORKDIR, exist_ok=True)

    doc = Document()

    # Create a 4x3 table
    table = doc.add_table(rows=4, cols=3)
    table.style = 'Table Grid'

    # Set table data
    data = [
        ['Feature', 'Basic', 'Premium'],
        ['Storage', '5 GB', '100 GB'],
        ['Support', 'Email', '24/7 Phone'],
        ['Price', '$5/mo', '$25/mo'],
    ]
    for i, row_data in enumerate(data):
        for j, text in enumerate(row_data):
            table.cell(i, j).text = text

    # Set table-level borders: all sides 0.5pt (sz=4) solid black, including inside
    set_table_borders(
        table,
        outer_val="single", outer_sz="4", outer_color="000000",
        inner_val="single", inner_sz="4", inner_color="000000",
    )

    # Also set cell-level borders for every cell to ensure consistent rendering
    black = ("single", "4", "000000")
    for row in table.rows:
        for cell in row.cells:
            set_cell_borders(cell, top=black, bottom=black, left=black, right=black)

    doc.save(OUTPUT)
    print(f'Initial file created: {OUTPUT}')

    # GUI-ready startup: open the file in LibreOffice Writer
    launch_gui(f'libreoffice --writer "{OUTPUT}"', delay_sec=2.0)
    print('GUI_READY: launched LibreOffice Writer with DISPLAY=:0')


create_initial()
