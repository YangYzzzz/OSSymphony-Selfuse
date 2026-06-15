"""
Reward Script: Fillable form fields in retainer agreement
Task ID: writer_legal_074
Domain: libreoffice_writer
Scoring:
  Component 1 (0.30): SDT content controls exist (>= 5)
  Component 2 (0.30): Fields have correct names/tags
  Component 3 (0.20): Bracketed placeholders removed
  Component 4 (0.20): All SDT elements are text-type
"""

import os
import re

WORKDIR = '/home/user'
TASK_ID = 'writer_legal_074'

# The 5 required field names from the task
REQUIRED_FIELDS = {'effective_date', 'client_name', 'client_address', 'retainer_amount', 'billing_rate'}

# The bracketed placeholders that should be removed
PLACEHOLDERS = ['[CLIENT_NAME]', '[CLIENT_ADDRESS]', '[RETAINER_AMOUNT]', '[BILLING_RATE]', '[EFFECTIVE_DATE]']


def verify_task(file_path):
    """
    Verify task completion with progressive scoring.
    Returns: float between 0.0 and 1.0
    """
    total_score = 0.0

    try:
        from docx import Document
        from lxml import etree
        doc = Document(file_path)
    except Exception as e:
        print(f"CRITICAL: Cannot load file {file_path}: {e}")
        print("REWARD: 0.0")
        return 0.0

    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    body = doc.element.body

    # Find all SDT (structured document tag) elements — these are content controls / form fields
    sdt_elements = body.findall('.//w:sdt', ns)
    num_sdts = len(sdt_elements)

    # Also check for legacy form fields (ffData) as an alternative implementation
    ffdata_elements = body.findall('.//w:ffData', ns)
    num_ffdata = len(ffdata_elements)

    # Collect SDT field info
    sdt_tags = []
    sdt_aliases = []
    sdt_is_text = []
    for sdt in sdt_elements:
        sdtPr = sdt.find('w:sdtPr', ns)
        if sdtPr is not None:
            tag_el = sdtPr.find('w:tag', ns)
            alias_el = sdtPr.find('w:alias', ns)
            text_el = sdtPr.find('w:text', ns)
            wval = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val'
            tag_val = tag_el.get(wval) if tag_el is not None else None
            alias_val = alias_el.get(wval) if alias_el is not None else None
            sdt_tags.append(tag_val)
            sdt_aliases.append(alias_val)
            sdt_is_text.append(text_el is not None)

    # Collect ffData field info (legacy form fields)
    ff_names = []
    ff_is_text = []
    for ff in ffdata_elements:
        name_el = ff.find('w:name', ns)
        text_input = ff.find('w:textInput', ns)
        wval = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val'
        name = name_el.get(wval) if name_el is not None else None
        ff_names.append(name)
        ff_is_text.append(text_input is not None)

    # Decide which field system is used
    using_sdt = num_sdts >= 5
    using_ffdata = num_ffdata >= 5
    has_fields = using_sdt or using_ffdata

    # Component 1: Form fields exist (>= 5 fields) (0.30 points)
    # This checks that the document has been converted from plain text placeholders to form fields
    try:
        if has_fields:
            field_count = max(num_sdts, num_ffdata)
            print(f"PASS: Component 1 — {field_count} form fields found (SDT={num_sdts}, ffData={num_ffdata}) (0.30 pts)")
            total_score += 0.30
        else:
            print(f"FAIL: Component 1 — Need >= 5 form fields, found SDT={num_sdts}, ffData={num_ffdata}")
    except Exception as e:
        print(f"ERROR: Component 1 — {e}")

    # Component 2: Fields have correct names/tags matching required fields (0.30 points)
    # Each of the 5 required field names must appear as a tag/alias/name in the form fields
    try:
        if using_sdt:
            # Use tags or aliases (whichever is populated)
            field_names_found = set()
            for tag, alias in zip(sdt_tags, sdt_aliases):
                if tag:
                    field_names_found.add(tag.lower().strip())
                if alias:
                    field_names_found.add(alias.lower().strip())
        elif using_ffdata:
            field_names_found = set()
            for name in ff_names:
                if name:
                    field_names_found.add(name.lower().strip())
        else:
            field_names_found = set()

        matched = REQUIRED_FIELDS & field_names_found
        match_ratio = len(matched) / len(REQUIRED_FIELDS)
        points = round(0.30 * match_ratio, 2)

        if match_ratio == 1.0:
            print(f"PASS: Component 2 — All 5 required field names found: {matched} (0.30 pts)")
            total_score += 0.30
        elif match_ratio > 0:
            print(f"PARTIAL: Component 2 — {len(matched)}/5 field names matched: {matched}, missing: {REQUIRED_FIELDS - matched} ({points} pts)")
            total_score += points
        else:
            print(f"FAIL: Component 2 — No required field names found. Fields present: {field_names_found}")
    except Exception as e:
        print(f"ERROR: Component 2 — {e}")

    # Component 3: Bracketed placeholders removed from document text (0.20 points)
    # In the golden file, all [PLACEHOLDER] text should be replaced by form fields
    try:
        full_text = '\n'.join(para.text for para in doc.paragraphs)
        placeholders_remaining = [p for p in PLACEHOLDERS if p in full_text]

        if len(placeholders_remaining) == 0:
            print(f"PASS: Component 3 — All bracketed placeholders removed (0.20 pts)")
            total_score += 0.20
        else:
            # Partial credit based on how many were removed
            removed_count = len(PLACEHOLDERS) - len(placeholders_remaining)
            ratio = removed_count / len(PLACEHOLDERS)
            points = round(0.20 * ratio, 2)
            print(f"PARTIAL: Component 3 — {len(placeholders_remaining)} placeholders still present: {placeholders_remaining} ({points} pts)")
            total_score += points
    except Exception as e:
        print(f"ERROR: Component 3 — {e}")

    # Component 4: All form field elements are text-type (0.20 points)
    # The task specifically asks for text input fields
    try:
        if has_fields:
            if using_sdt:
                text_field_checks = sdt_is_text
            else:
                text_field_checks = ff_is_text

            text_count = sum(1 for t in text_field_checks if t)
            total_fields = len(text_field_checks)

            if total_fields > 0 and text_count == total_fields:
                print(f"PASS: Component 4 — All {total_fields} fields are text-type (0.20 pts)")
                total_score += 0.20
            elif text_count > 0:
                ratio = text_count / total_fields
                points = round(0.20 * ratio, 2)
                print(f"PARTIAL: Component 4 — {text_count}/{total_fields} fields are text-type ({points} pts)")
                total_score += points
            else:
                print(f"FAIL: Component 4 — No text-type fields found among {total_fields} fields")
        else:
            print(f"FAIL: Component 4 — No form fields to check type on")
    except Exception as e:
        print(f"ERROR: Component 4 — {e}")

    final_score = min(round(total_score, 2), 1.0)
    print(f"\nScore: {total_score}/1.0")
    print(f"REWARD: {final_score}")
    return final_score


# Default: test against canonical artifact path
file_path = f'{WORKDIR}/{TASK_ID}.docx'
if not os.path.exists(file_path):
    print(f"File not found: {file_path}")
    print("REWARD: 0.0")
else:
    verify_task(file_path)
