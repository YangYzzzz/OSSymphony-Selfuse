"""
Initial Setup: Table with default 0.5pt solid black borders (pre-task state)
Task ID: writer_tbl_049
Domain: libreoffice_writer

Creates a 4x3 table with 0.5pt solid black borders on all sides.
The agent task is to apply custom border styles.
"""

import os
import shlex
import subprocess
import time
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

WORKDIR = '/home/user'
TASK_ID = 'writer_tbl_049'
OUTPUT = f'{WORKDIR}/{TASK_ID}.docx'


def make_border_element(val, sz, space, color):
    """Create a border XML element with given properties."""
    border = OxmlElement('w:' + 'dummy')  # placeholder
    border.set(qn('w:val'), val)
    border.set(qn('w:sz'), str(sz))
    border.set(qn('w:space'), str(space))
    border.set(qn('w:color'), color)
    return border


def set_table_borders(table, top_val, top_sz, bottom_val, bottom_sz,
                       left_val, left_sz, right_val, right_sz,
                       insideH_val, insideH_sz, insideV_val, insideV_sz,
                       color='000000', dashed_color='000000'):
    """
    Set table-level borders using XML manipulation.
    sz is in eighths of a point (e.g. 4 = 0.5pt, 16 = 2pt).
    """
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)

    # Remove existing tblBorders if present
    for existing in tblPr.findall(qn('w:tblBorders')):
        tblPr.remove(existing)

    tblBorders = OxmlElement('w:tblBorders')

    def make_border(tag, val, sz, bcolor='000000'):
        el = OxmlElement(tag)
        el.set(qn('w:val'), val)
        el.set(qn('w:sz'), str(sz))
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), bcolor)
        return el

    if top_val:
        tblBorders.append(make_border('w:top', top_val, top_sz, color))
    if left_val:
        tblBorders.append(make_border('w:left', left_val, left_sz, color))
    if bottom_val:
        tblBorders.append(make_border('w:bottom', bottom_val, bottom_sz, color))
    if right_val:
        tblBorders.append(make_border('w:right', right_val, right_sz, color))
    if insideH_val:
        tblBorders.append(make_border('w:insideH', insideH_val, insideH_sz, color))
    if insideV_val:
        tblBorders.append(make_border('w:insideV', insideV_val, insideV_sz, color))

    tblPr.append(tblBorders)


def clear_cell_borders(table):
    """Clear per-cell border overrides so table-level borders apply."""
    for row in table.rows:
        for cell in row.cells:
            tc = cell._tc
            tcPr = tc.tcPr
            if tcPr is not None:
                for existing in tcPr.findall(qn('w:tcBorders')):
                    tcPr.remove(existing)


def launch_gui(command: str, delay_sec: float = 1.0):
    """Launch GUI app on VM display without blocking script exit."""
    env = os.environ.copy()
    env["DISPLAY"] = ":0"
    subprocess.Popen(
        shlex.split(command),
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        env=env,
    )
    time.sleep(delay_sec)


def create_initial():
    doc = Document()

    # Create 4x3 table
    table = doc.add_table(rows=4, cols=3)
    table.style = 'Table Grid'

    # Populate cells
    data = [
        ['Rank', 'Name', 'Score'],
        ['1', 'Maxwell', '98'],
        ['2', 'Rivera', '95'],
        ['3', 'Thompson', '92'],
    ]
    for r_idx, row_data in enumerate(data):
        for c_idx, text in enumerate(row_data):
            table.cell(r_idx, c_idx).text = text

    # Set default 0.5pt solid black borders on all sides (sz=4 = 0.5pt in 1/8pt units)
    set_table_borders(
        table,
        top_val='single', top_sz=4,
        bottom_val='single', bottom_sz=4,
        left_val='single', left_sz=4,
        right_val='single', right_sz=4,
        insideH_val='single', insideH_sz=4,
        insideV_val='single', insideV_sz=4,
        color='000000',
    )

    # Clear any cell-level border overrides
    clear_cell_borders(table)

    doc.save(OUTPUT)
    print(f'Initial file created: {OUTPUT}')

    # GUI-ready startup: open the document in LibreOffice Writer
    launch_gui(f'libreoffice --writer "{OUTPUT}"', delay_sec=2.0)
    print('GUI_READY: launched LibreOffice Writer with DISPLAY=:0')


create_initial()
