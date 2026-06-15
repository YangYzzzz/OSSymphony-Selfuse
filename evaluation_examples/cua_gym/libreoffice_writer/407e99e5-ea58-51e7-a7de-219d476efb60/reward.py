"""
Reward Script: Mail merge field insertion for client notification letter
Task ID: writer_legal_041
Domain: libreoffice_writer
Scoring:
  Component 1 (0.25): client_name MERGEFIELD present
  Component 2 (0.25): client_address MERGEFIELD present
  Component 3 (0.25): case_number MERGEFIELD present
  Component 4 (0.15): hearing_date MERGEFIELD present
  Component 5 (0.10): No leftover bracket placeholders remain
"""

import os

WORKDIR = '/home/user'
TASK_ID = 'writer_legal_041'


def extract_merge_field_names(doc):
    """Extract all MERGEFIELD names from the document XML."""
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    merge_fields = []
    for instr in doc.element.findall('.//w:instrText', ns):
        text = instr.text.strip() if instr.text else ''
        if 'MERGEFIELD' in text:
            field_name = text.replace('MERGEFIELD', '').strip()
            merge_fields.append(field_name)
    return merge_fields


def has_leftover_placeholders(doc):
    """Check if any bracket placeholders like [CLIENT_NAME] remain."""
    placeholders = ['[CLIENT_NAME]', '[CLIENT_ADDRESS]', '[CASE_NUMBER]', '[HEARING_DATE]']
    full_text = '\n'.join(p.text for p in doc.paragraphs)
    found = []
    for ph in placeholders:
        if ph in full_text:
            found.append(ph)
    return found


def verify_task(file_path):
    """
    Verify task completion with progressive scoring.
    Returns: float between 0.0 and 1.0
    """
    from docx import Document

    total_score = 0.0

    try:
        doc = Document(file_path)
    except Exception as e:
        print(f"CRITICAL: Cannot load file {file_path}: {e}")
        print("REWARD: 0.0")
        return 0.0

    # Extract all merge fields from document
    try:
        merge_fields = extract_merge_field_names(doc)
        print(f"INFO: Found {len(merge_fields)} MERGEFIELD(s): {merge_fields}")
    except Exception as e:
        print(f"ERROR: Could not extract merge fields: {e}")
        merge_fields = []

    # Component 1: client_name MERGEFIELD present (0.25 points)
    try:
        client_name_count = sum(1 for f in merge_fields if f == 'client_name')
        if client_name_count >= 1:
            print(f"PASS: Component 1 — client_name MERGEFIELD found ({client_name_count} instance(s)) (0.25 pts)")
            total_score += 0.25
        else:
            print(f"FAIL: Component 1 — client_name MERGEFIELD not found")
    except Exception as e:
        print(f"ERROR: Component 1 — {e}")

    # Component 2: client_address MERGEFIELD present (0.25 points)
    try:
        client_address_count = sum(1 for f in merge_fields if f == 'client_address')
        if client_address_count >= 1:
            print(f"PASS: Component 2 — client_address MERGEFIELD found ({client_address_count} instance(s)) (0.25 pts)")
            total_score += 0.25
        else:
            print(f"FAIL: Component 2 — client_address MERGEFIELD not found")
    except Exception as e:
        print(f"ERROR: Component 2 — {e}")

    # Component 3: case_number MERGEFIELD present (0.25 points)
    try:
        case_number_count = sum(1 for f in merge_fields if f == 'case_number')
        if case_number_count >= 1:
            print(f"PASS: Component 3 — case_number MERGEFIELD found ({case_number_count} instance(s)) (0.25 pts)")
            total_score += 0.25
        else:
            print(f"FAIL: Component 3 — case_number MERGEFIELD not found")
    except Exception as e:
        print(f"ERROR: Component 3 — {e}")

    # Component 4: hearing_date MERGEFIELD present (0.15 points)
    try:
        hearing_date_count = sum(1 for f in merge_fields if f == 'hearing_date')
        if hearing_date_count >= 1:
            print(f"PASS: Component 4 — hearing_date MERGEFIELD found ({hearing_date_count} instance(s)) (0.15 pts)")
            total_score += 0.15
        else:
            print(f"FAIL: Component 4 — hearing_date MERGEFIELD not found")
    except Exception as e:
        print(f"ERROR: Component 4 — {e}")

    # Component 5: No leftover bracket placeholders remain (0.10 points)
    # This component ONLY awards points if merge fields ARE present (to avoid scoring initial state)
    try:
        leftover = has_leftover_placeholders(doc)
        if len(merge_fields) > 0 and len(leftover) == 0:
            print(f"PASS: Component 5 — No leftover placeholders found (0.10 pts)")
            total_score += 0.10
        elif len(merge_fields) == 0:
            print(f"FAIL: Component 5 — No merge fields present, skipping placeholder check")
        else:
            print(f"FAIL: Component 5 — Leftover placeholders found: {leftover}")
    except Exception as e:
        print(f"ERROR: Component 5 — {e}")

    final_score = min(total_score, 1.0)
    print(f"\nScore: {total_score}/1.0")
    print(f"REWARD: {final_score}")
    return final_score


# Entry point
file_path = f'{WORKDIR}/{TASK_ID}.docx'
if not os.path.exists(file_path):
    print(f"File not found: {file_path}")
    print("REWARD: 0.0")
else:
    verify_task(file_path)
