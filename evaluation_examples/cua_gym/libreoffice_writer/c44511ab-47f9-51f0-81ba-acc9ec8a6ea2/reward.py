"""
FINAL REWARD SCRIPT - SUCCESS
Task: My draft is peppered with the exact string "TODO:" and I want each of those placeholders to scream at me during proofreading. In LibreOffice Writer, how can I instantly locate every occurrence of "TODO:" and apply direct formatting—specifically bold plus the font color set to #FF0000—without stepping through them one by one?
Generated: 2025-09-10 15:23:51
Status: success
Model: azure-o3
Total Steps: 3
"""

import os
from docx import Document
from docx.shared import RGBColor

"""
Reward Script: Verify that every occurrence of the exact string "TODO:" in the provided
DOCX file is formatted in BOTH bold and font-color #FF0000.
Progressive scoring (0.0 – 1.0):
  • 0.30 if at least one "TODO:" placeholder is found (task detected)
  • +0.35 if ALL placeholders are bold
  • +0.35 if ALL placeholders are red (#FF0000)
Returns exactly 1.0 when every occurrence satisfies both formatting rules.
"""

FILE_PATH = "/home/user/my_draft_is_peppered_with_the_exact_string_todo_and_i_want_each_of_those_placeholders_to_scream_at_m.docx"
TARGET = "TODO:"


def _is_rgb_red(run):
    """Return True if run font color is exactly #FF0000."""
    try:
        color = run.font.color
        return bool(color and color.rgb and color.rgb == RGBColor(0xFF, 0x00, 0x00))
    except Exception:
        return False


def verify_todo_formatting(file_path: str) -> float:
    print(f"Checking document: {file_path}")

    # ---------- Safety checks ----------
    if not os.path.exists(file_path):
        print("✗ File not found – cannot evaluate task")
        print("REWARD: 0.0")
        return 0.0
    try:
        doc = Document(file_path)
    except Exception as e:
        print(f"✗ Failed to load DOCX: {e}")
        print("REWARD: 0.0")
        return 0.0

    # ---------- Scan document ----------
    total_occ = 0        # total "TODO:" strings discovered
    bold_correct = 0     # occurrences fully bold
    red_correct = 0      # occurrences fully red

    for para in doc.paragraphs:
        # Build a character-level list: (char, is_bold, is_red)
        char_info = []
        for run in para.runs:
            bold_flag = bool(run.bold)
            red_flag = _is_rgb_red(run)
            for ch in run.text:
                char_info.append((ch, bold_flag, red_flag))

        if not char_info:
            continue

        para_text = ''.join(ch for ch, _, _ in char_info)
        idx = 0
        while True:
            idx = para_text.find(TARGET, idx)
            if idx == -1:
                break
            segment = char_info[idx: idx + len(TARGET)]
            if len(segment) == len(TARGET):
                total_occ += 1
                if all(bold for _, bold, _ in segment):
                    bold_correct += 1
                if all(red for _, _, red in segment):
                    red_correct += 1
            idx += len(TARGET)

    # ---------- Reporting ----------
    print(f"Total 'TODO:' occurrences: {total_occ}")
    print(f"Correctly bold: {bold_correct}/{total_occ}")
    print(f"Correctly colored red: {red_correct}/{total_occ}")

    # ---------- Scoring ----------
    score = 0.0
    if total_occ > 0:
        score += 0.30                       # Found at least one placeholder
        if bold_correct == total_occ:
            score += 0.35                  # All bold
        if red_correct == total_occ:
            score += 0.35                  # All red

    # Avoid floating point artefacts and cap at 1.0
    final_score = min(round(score, 4), 1.0)
    print(f"REWARD: {final_score}")
    return final_score


# -------------- Execute verification when run directly --------------
if __name__ == "__main__":
    verify_todo_formatting(FILE_PATH)

