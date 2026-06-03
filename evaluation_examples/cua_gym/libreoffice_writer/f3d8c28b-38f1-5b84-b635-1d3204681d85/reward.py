"""
Reward Script: Insert a text frame with pull quote, blue border, light blue background
Task ID: writer_rd_023
Domain: libreoffice_writer
Scoring:
  Component 1: Text frame exists with correct width ~7cm (0.20)
  Component 2: Horizontal centering (0.15)
  Component 3: Border color #003366 and ~2pt width (0.20)
  Component 4: Background fill #E6F0FF (0.15)
  Component 5: Pull quote text present, italic, 14pt, center-aligned (0.20)
  Component 6: Text wrapping on both sides with ~0.4cm spacing (0.10)
"""

import os
from docx import Document
from docx.oxml.ns import qn
import lxml.etree as etree

WORKDIR = '/home/user'
TASK_ID = 'writer_rd_023'

NS = {
    'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
    'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'wps': 'http://schemas.microsoft.com/office/word/2010/wordprocessingShape',
    'v': 'urn:schemas-microsoft-com:vml',
}


def find_text_frame(doc):
    """Find the first text box drawing element in the document body."""
    body = doc.element.body
    drawings = body.findall('.//w:drawing', NS)
    for d in drawings:
        # Look for wps:wsp with txBox attribute
        wsps = d.findall('.//wps:wsp', NS)
        for wsp in wsps:
            cnv = wsp.find('wps:cNvSpPr', NS)
            if cnv is not None and cnv.get('txBox') == '1':
                return d, wsp
    return None, None


def verify_task(file_path):
    """
    Verify task completion with progressive scoring.
    Returns: float between 0.0 and 1.0
    """
    total_score = 0.0

    try:
        doc = Document(file_path)
    except Exception as e:
        print(f"CRITICAL: Cannot load file {file_path}: {e}")
        print("REWARD: 0.0")
        return 0.0

    drawing, wsp = find_text_frame(doc)

    if drawing is None or wsp is None:
        print("FAIL: No text frame (text box) found in document")
        print("REWARD: 0.0")
        return 0.0

    # Component 1: Text frame width ~7cm (2520000 EMU, tolerance ±5%) (0.20 points)
    try:
        # Check extent cx
        anchor = drawing.find('.//wp:anchor', NS)
        if anchor is None:
            anchor = drawing.find('.//wp:inline', NS)
        extent = anchor.find('wp:extent', NS) if anchor is not None else None
        if extent is not None:
            cx = int(extent.get('cx', '0'))
            # 7cm = 2520000 EMU; allow ±5% tolerance
            target_cx = 2520000
            if abs(cx - target_cx) / target_cx <= 0.05:
                print(f"PASS: Component 1 — Frame width {cx} EMU (~{cx/360000:.1f} cm), close to 7cm (0.20 pts)")
                total_score += 0.20
            else:
                print(f"FAIL: Component 1 — Frame width {cx} EMU (~{cx/360000:.1f} cm), expected ~2520000 EMU (7cm)")
        else:
            print("FAIL: Component 1 — No extent element found")
    except Exception as e:
        print(f"ERROR: Component 1 — {e}")

    # Component 2: Horizontal centering (0.15 points)
    try:
        posH = anchor.find('wp:positionH', NS) if anchor is not None else None
        if posH is not None:
            rel_from = posH.get('relativeFrom', '')
            align_elem = posH.find('wp:align', NS)
            if align_elem is not None and align_elem.text == 'center':
                print(f"PASS: Component 2 — Horizontally centered (relativeFrom={rel_from}) (0.15 pts)")
                total_score += 0.15
            else:
                align_text = align_elem.text if align_elem is not None else 'N/A'
                print(f"FAIL: Component 2 — Horizontal align is '{align_text}', expected 'center'")
        else:
            print("FAIL: Component 2 — No positionH element found")
    except Exception as e:
        print(f"ERROR: Component 2 — {e}")

    # Component 3: Border color #003366 and ~2pt line width (0.20 points)
    try:
        spPr = wsp.find('wps:spPr', NS)
        ln = spPr.find('a:ln', NS) if spPr is not None else None
        if ln is not None:
            ln_w = int(ln.get('w', '0'))
            # 2pt = 25400 EMU; allow ±20% tolerance
            ln_fill = ln.find('a:solidFill', NS)
            ln_color_elem = ln_fill.find('a:srgbClr', NS) if ln_fill is not None else None
            ln_color = ln_color_elem.get('val', '').upper() if ln_color_elem is not None else ''

            width_ok = abs(ln_w - 25400) / 25400 <= 0.20 if ln_w > 0 else False
            color_ok = ln_color == '003366'

            if width_ok and color_ok:
                print(f"PASS: Component 3 — Border: {ln_w} EMU (~{ln_w/12700:.1f}pt), color #{ln_color} (0.20 pts)")
                total_score += 0.20
            else:
                print(f"FAIL: Component 3 — Border width={ln_w} EMU (ok={width_ok}), color=#{ln_color} (ok={color_ok})")
        else:
            print("FAIL: Component 3 — No line element found in shape properties")
    except Exception as e:
        print(f"ERROR: Component 3 — {e}")

    # Component 4: Background fill #E6F0FF (0.15 points)
    try:
        spPr = wsp.find('wps:spPr', NS)
        solid_fill = spPr.find('a:solidFill', NS) if spPr is not None else None
        if solid_fill is not None:
            clr_elem = solid_fill.find('a:srgbClr', NS)
            bg_color = clr_elem.get('val', '').upper() if clr_elem is not None else ''
            if bg_color == 'E6F0FF':
                print(f"PASS: Component 4 — Background fill #{bg_color} (0.15 pts)")
                total_score += 0.15
            else:
                print(f"FAIL: Component 4 — Background fill is #{bg_color}, expected #E6F0FF")
        else:
            print("FAIL: Component 4 — No solidFill found in shape properties")
    except Exception as e:
        print(f"ERROR: Component 4 — {e}")

    # Component 5: Pull quote text present, italic, 14pt, center-aligned (0.20 points)
    try:
        txbx = wsp.find('.//wps:txbx', NS)
        txbx_content = txbx.find('w:txbxContent', NS) if txbx is not None else None
        if txbx_content is not None:
            # Get all text from txbxContent
            all_text = ''
            for p_elem in txbx_content.findall('w:p', NS):
                for r_elem in p_elem.findall('w:r', NS):
                    t_elem = r_elem.find('w:t', NS)
                    if t_elem is not None and t_elem.text:
                        all_text += t_elem.text

            expected_quote = 'Innovation is not just about technology; it is about solving real problems.'
            text_ok = expected_quote.lower() in all_text.lower()

            # Check italic and size on runs
            italic_ok = False
            size_ok = False
            center_ok = False

            for p_elem in txbx_content.findall('w:p', NS):
                # Check paragraph center alignment
                pPr = p_elem.find('w:pPr', NS)
                if pPr is not None:
                    jc = pPr.find('w:jc', NS)
                    if jc is not None and jc.get(qn('w:val')) == 'center':
                        center_ok = True

                for r_elem in p_elem.findall('w:r', NS):
                    rPr = r_elem.find('w:rPr', NS)
                    if rPr is not None:
                        i_elem = rPr.find('w:i', NS)
                        if i_elem is not None:
                            # w:i without val means True; val="true"/"1" also True
                            val = i_elem.get(qn('w:val'))
                            if val is None or val in ('true', '1'):
                                italic_ok = True
                        sz_elem = rPr.find('w:sz', NS)
                        if sz_elem is not None:
                            sz_val = int(sz_elem.get(qn('w:val'), '0'))
                            # 14pt = 28 half-points
                            if sz_val == 28:
                                size_ok = True

            sub_score = 0.0
            if text_ok:
                sub_score += 0.08
            if italic_ok:
                sub_score += 0.04
            if size_ok:
                sub_score += 0.04
            if center_ok:
                sub_score += 0.04

            if sub_score > 0:
                print(f"PASS: Component 5 — text={text_ok}, italic={italic_ok}, size14pt={size_ok}, center={center_ok} ({sub_score:.2f} pts)")
                total_score += sub_score
            else:
                print(f"FAIL: Component 5 — text={text_ok}, italic={italic_ok}, size14pt={size_ok}, center={center_ok}")
        else:
            print("FAIL: Component 5 — No txbxContent found")
    except Exception as e:
        print(f"ERROR: Component 5 — {e}")

    # Component 6: Text wrapping on both sides with ~0.4cm spacing (0.10 points)
    try:
        if anchor is not None:
            wrap_square = anchor.find('wp:wrapSquare', NS)
            wrap_tight = anchor.find('wp:wrapTight', NS)
            wrap_through = anchor.find('wp:wrapThrough', NS)

            if wrap_square is not None:
                wrap_text = wrap_square.get('wrapText', '')
                # Check spacing via distT/distB/distL/distR on anchor
                distL = int(anchor.get('distL', '0'))
                distR = int(anchor.get('distR', '0'))
                # 0.4cm = 144000 EMU; allow ±30% tolerance
                target_dist = 144000
                both_sides = wrap_text == 'bothSides'
                spacing_ok = (abs(distL - target_dist) / target_dist <= 0.30 and
                              abs(distR - target_dist) / target_dist <= 0.30) if target_dist > 0 else False

                if both_sides and spacing_ok:
                    print(f"PASS: Component 6 — wrapSquare bothSides, distL={distL}, distR={distR} (0.10 pts)")
                    total_score += 0.10
                elif both_sides:
                    print(f"PARTIAL: Component 6 — wrapSquare bothSides OK, spacing off (distL={distL}, distR={distR}) (0.05 pts)")
                    total_score += 0.05
                else:
                    print(f"FAIL: Component 6 — wrapText='{wrap_text}', expected 'bothSides'")
            elif wrap_tight is not None or wrap_through is not None:
                print("PARTIAL: Component 6 — wrap exists but not wrapSquare (0.03 pts)")
                total_score += 0.03
            else:
                print("FAIL: Component 6 — No wrap element found")
        else:
            print("FAIL: Component 6 — No anchor element")
    except Exception as e:
        print(f"ERROR: Component 6 — {e}")

    final_score = min(round(total_score, 2), 1.0)
    print(f"\nScore: {total_score:.2f}/1.0")
    print(f"REWARD: {final_score}")
    return final_score


# Persistence hook for LibreOffice Writer
def persist_app_state():
    import time
    os.environ["DISPLAY"] = ":0"
    try:
        import pyautogui
        pyautogui.hotkey("ctrl", "s")
        time.sleep(0.8)
        print("PERSIST: ctrl+s sent for libreoffice_writer")
    except Exception as e:
        print(f"PERSIST_WARN: save hook failed: {e}")


# Entry point
file_path = f'{WORKDIR}/{TASK_ID}.docx'
persist_app_state()

import time
time.sleep(0.5)

if not os.path.exists(file_path):
    print(f"File not found: {file_path}")
    print("REWARD: 0.0")
else:
    verify_task(file_path)
