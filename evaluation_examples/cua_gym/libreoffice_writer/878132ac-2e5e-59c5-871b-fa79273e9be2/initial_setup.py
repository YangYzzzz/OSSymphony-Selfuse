"""
Initial Setup: Set column widths in a table with test scores
Task ID: writer_tbl_017
Domain: libreoffice_writer
"""

import os
import shlex
import subprocess
import time
from docx import Document
from docx.shared import Pt

WORKDIR = '/home/user/Desktop'  # Task context places file on ~/Desktop/
TASK_ID = 'test_scores'
OUTPUT = f'{WORKDIR}/{TASK_ID}.docx'


def launch_gui(command: str, delay_sec: float = 1.0):
    """Launch GUI app on VM display without blocking script exit."""
    env = os.environ.copy()
    env["DISPLAY"] = ":0"
    subprocess.Popen(
        shlex.split(command),
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        env=env,
    )
    time.sleep(delay_sec)


def create_initial():
    # Ensure Desktop directory exists
    os.makedirs(WORKDIR, exist_ok=True)

    doc = Document()

    # Create a 4-row x 2-column table with automatic column widths
    table = doc.add_table(rows=4, cols=2)
    table.style = 'Table Grid'

    # Row 1 (header): Name | Score
    table.cell(0, 0).text = 'Name'
    table.cell(0, 1).text = 'Score'

    # Row 2
    table.cell(1, 0).text = 'Oliver'
    table.cell(1, 1).text = '87'

    # Row 3
    table.cell(2, 0).text = 'Sophie'
    table.cell(2, 1).text = '93'

    # Row 4
    table.cell(3, 0).text = 'Liam'
    table.cell(3, 1).text = '78'

    # NOTE: Column widths are intentionally left as automatic (not set)
    # The task asks the agent to set them to 5 cm and 3 cm respectively

    doc.save(OUTPUT)
    print(f'Initial file created: {OUTPUT}')

    # GUI-ready startup: open the file in LibreOffice Writer
    launch_gui(f'libreoffice --writer "{OUTPUT}"', delay_sec=2.0)
    print('GUI_READY: launched LibreOffice Writer with DISPLAY=:0')


create_initial()
