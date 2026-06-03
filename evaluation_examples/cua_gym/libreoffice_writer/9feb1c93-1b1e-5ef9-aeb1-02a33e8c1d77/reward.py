"""
Reward Script: Mail merge field replacement in client letter template
Task ID: writer_biz_057
Domain: libreoffice_writer
Scoring:
  Component 1 (0.30): Address block placeholders replaced with merge fields (paras 6-8)
  Component 2 (0.25): Salutation placeholder replaced with merge field (para 9)
  Component 3 (0.25): Body placeholder replaced with merge field (para 13)
  Component 4 (0.20): All merge fields use correct field names (Name, Company, Address)
"""

import os
from docx import Document
from lxml import etree

WORKDIR = '/home/user'
TASK_ID = 'writer_biz_057'

WNS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NS = {'w': WNS}


def extract_merge_fields_from_paragraph(para):
    """Extract MERGEFIELD names from a paragraph's XML by finding instrText elements."""
    fields = []
    for instr in para._element.iter(f'{{{WNS}}}instrText'):
        text = (instr.text or '').strip()
        if text.startswith('MERGEFIELD'):
            field_name = text.replace('MERGEFIELD', '').strip()
            fields.append(field_name)
    return fields


def has_fld_char_structure(para):
    """Check if paragraph contains fldChar begin/separate/end structure (real merge fields)."""
    fld_types = []
    for fc in para._element.iter(f'{{{WNS}}}fldChar'):
        fld_type = fc.get(f'{{{WNS}}}fldCharType')
        if fld_type:
            fld_types.append(fld_type)
    return 'begin' in fld_types and 'separate' in fld_types and 'end' in fld_types


def paragraph_has_placeholder(para, placeholder):
    """Check if paragraph still contains the original placeholder text."""
    return placeholder in para.text


def verify_task(file_path):
    """
    Verify task completion with progressive scoring.
    Returns: float between 0.0 and 1.0
    """
    total_score = 0.0

    try:
        doc = Document(file_path)
    except Exception as e:
        print(f"CRITICAL: Cannot load file {file_path}: {e}")
        print("REWARD: 0.0")
        return 0.0

    paras = doc.paragraphs

    # Sanity check: document should have enough paragraphs
    if len(paras) < 14:
        print(f"CRITICAL: Expected at least 14 paragraphs, found {len(paras)}")
        print("REWARD: 0.0")
        return 0.0

    # Component 1: Address block placeholders replaced with merge fields (0.30 pts)
    # Paras 6, 7, 8 should have MERGEFIELD Name, Company, Address respectively
    # and should NOT contain [Client Name], [Company], [Address] placeholders
    try:
        addr_score = 0.0
        addr_checks = [
            (6, '[Client Name]', 'Name'),
            (7, '[Company]', 'Company'),
            (8, '[Address]', 'Address'),
        ]
        for para_idx, old_placeholder, expected_field in addr_checks:
            para = paras[para_idx]
            merge_fields = extract_merge_fields_from_paragraph(para)
            has_fld = has_fld_char_structure(para)
            still_has_placeholder = paragraph_has_placeholder(para, old_placeholder)

            if expected_field in merge_fields and has_fld and not still_has_placeholder:
                addr_score += 0.10
                print(f"  PASS: Para {para_idx} has MERGEFIELD {expected_field}")
            else:
                print(f"  FAIL: Para {para_idx} — merge_fields={merge_fields}, has_fld={has_fld}, placeholder_present={still_has_placeholder}")

        if addr_score > 0:
            print(f"PASS: Component 1 — Address block merge fields ({addr_score:.2f} pts)")
            total_score += addr_score
        else:
            print(f"FAIL: Component 1 — No address block merge fields found")
    except Exception as e:
        print(f"ERROR: Component 1 — {e}")

    # Component 2: Salutation placeholder replaced with merge field (0.25 pts)
    # Para 9: "Dear [Client Name]," -> "Dear «Name»,"
    try:
        para = paras[9]
        merge_fields = extract_merge_fields_from_paragraph(para)
        has_fld = has_fld_char_structure(para)
        still_has_placeholder = paragraph_has_placeholder(para, '[Client Name]')

        if 'Name' in merge_fields and has_fld and not still_has_placeholder:
            # Also verify "Dear" is still present
            if 'Dear' in para.text:
                print(f"PASS: Component 2 — Salutation has MERGEFIELD Name (0.25 pts)")
                total_score += 0.25
            else:
                print(f"FAIL: Component 2 — Merge field present but 'Dear' prefix missing")
        else:
            print(f"FAIL: Component 2 — merge_fields={merge_fields}, has_fld={has_fld}, placeholder_present={still_has_placeholder}")
    except Exception as e:
        print(f"ERROR: Component 2 — {e}")

    # Component 3: Body paragraph placeholder replaced with merge field (0.25 pts)
    # Para 13: "...collaboration with [Company] and..." -> "...collaboration with «Company» and..."
    try:
        para = paras[13]
        merge_fields = extract_merge_fields_from_paragraph(para)
        has_fld = has_fld_char_structure(para)
        still_has_placeholder = paragraph_has_placeholder(para, '[Company]')

        if 'Company' in merge_fields and has_fld and not still_has_placeholder:
            # Verify surrounding text is intact
            if 'collaboration' in para.text.lower():
                print(f"PASS: Component 3 — Body paragraph has MERGEFIELD Company (0.25 pts)")
                total_score += 0.25
            else:
                print(f"FAIL: Component 3 — Merge field present but surrounding text damaged")
        else:
            print(f"FAIL: Component 3 — merge_fields={merge_fields}, has_fld={has_fld}, placeholder_present={still_has_placeholder}")
    except Exception as e:
        print(f"ERROR: Component 3 — {e}")

    # Component 4: All merge fields use correct field names (0.20 pts)
    # Count total MERGEFIELD instrText occurrences across entire document
    # Expected: Name (x2), Company (x2), Address (x1) = 5 total
    try:
        all_fields = []
        for para in paras:
            all_fields.extend(extract_merge_fields_from_paragraph(para))

        name_count = all_fields.count('Name')
        company_count = all_fields.count('Company')
        address_count = all_fields.count('Address')

        # Must have at least the expected fields
        if name_count >= 2 and company_count >= 2 and address_count >= 1:
            print(f"PASS: Component 4 — All merge field names correct: Name={name_count}, Company={company_count}, Address={address_count} (0.20 pts)")
            total_score += 0.20
        elif name_count >= 1 and company_count >= 1 and address_count >= 1:
            # Partial: at least one of each field exists
            print(f"PARTIAL: Component 4 — Some merge fields present: Name={name_count}, Company={company_count}, Address={address_count} (0.10 pts)")
            total_score += 0.10
        else:
            print(f"FAIL: Component 4 — Missing merge fields: Name={name_count}, Company={company_count}, Address={address_count}")
    except Exception as e:
        print(f"ERROR: Component 4 — {e}")

    final_score = min(total_score, 1.0)
    print(f"\nScore: {total_score}/1.0")
    print(f"REWARD: {final_score}")
    return final_score


# Default: test against canonical artifact path
file_path = f'{WORKDIR}/{TASK_ID}.docx'
if not os.path.exists(file_path):
    print(f"File not found: {file_path}")
    print("REWARD: 0.0")
else:
    verify_task(file_path)
