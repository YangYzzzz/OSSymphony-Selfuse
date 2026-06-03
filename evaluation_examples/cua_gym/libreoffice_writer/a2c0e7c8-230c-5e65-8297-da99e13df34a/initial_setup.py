"""
Initial Setup: Revenue table with plain (unformatted) numbers in Revenue column
Task ID: writer_tbl_056
Domain: libreoffice_writer
"""

import os
import shlex
import subprocess
import time
from docx import Document
from docx.shared import Pt

WORKDIR = '/home/user/Desktop'  # VM path — file placed on Desktop as per task
TASK_ID = 'writer_tbl_056'
OUTPUT = f'{WORKDIR}/revenue_table.docx'


def launch_gui(command: str, delay_sec: float = 1.0):
    """Launch GUI app on VM display without blocking script exit."""
    env = os.environ.copy()
    env["DISPLAY"] = ":0"
    subprocess.Popen(
        shlex.split(command),
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        env=env,
    )
    time.sleep(delay_sec)


def create_initial():
    # Ensure Desktop directory exists
    os.makedirs(WORKDIR, exist_ok=True)

    doc = Document()

    # Add a table with 5 rows and 3 columns
    table = doc.add_table(rows=5, cols=3)
    table.style = 'Table Grid'

    # Row 1: Header
    table.cell(0, 0).text = 'Quarter'
    table.cell(0, 1).text = 'Revenue'
    table.cell(0, 2).text = 'Growth'

    # Row 2: Q1
    table.cell(1, 0).text = 'Q1'
    table.cell(1, 1).text = '45000'
    table.cell(1, 2).text = '5%'

    # Row 3: Q2
    table.cell(2, 0).text = 'Q2'
    table.cell(2, 1).text = '52000'
    table.cell(2, 2).text = '15%'

    # Row 4: Q3
    table.cell(3, 0).text = 'Q3'
    table.cell(3, 1).text = '48000'
    table.cell(3, 2).text = '-8%'

    # Row 5: Q4
    table.cell(4, 0).text = 'Q4'
    table.cell(4, 1).text = '61000'
    table.cell(4, 2).text = '27%'

    doc.save(OUTPUT)
    print(f'Initial file created: {OUTPUT}')

    # GUI-ready startup: open the file in LibreOffice Writer
    launch_gui(f'libreoffice --writer "{OUTPUT}"', delay_sec=2.0)
    print('GUI_READY: launched LibreOffice Writer with DISPLAY=:0')


create_initial()
